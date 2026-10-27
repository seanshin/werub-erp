"""렌더링 엔진 — Excel(openpyxl)·PDF 보안 fetcher (design §6.7/§11.0)."""
import io
from datetime import date

import pytest
from openpyxl import load_workbook

from erp.common.excel import SYNC_ROW_LIMIT, build_xlsx
from erp.common.pdf import BlockedUrlError, render_pdf, secure_url_fetcher


# ── Excel ───────────────────────────────────────────────────
def test_xlsx_roundtrip():
    content = build_xlsx(sheet_name="직원", headers=["사번", "이름", "입사일", "마약류"],
                         rows=[["E0001", "김민준", date(2020, 1, 1), True]])
    wb = load_workbook(io.BytesIO(content))
    ws = wb["직원"]
    assert [c.value for c in ws[1]] == ["사번", "이름", "입사일", "마약류"]
    assert [c.value for c in ws[2]] == ["E0001", "김민준", "2020-01-01", "Y"]


def test_xlsx_sync_limit_enforced():
    """5천 행 초과 — 잘림 대신 명시적 거부 (no silent caps)."""
    with pytest.raises(ValueError):
        build_xlsx(sheet_name="big", headers=["a"],
                   rows=[["x"]] * (SYNC_ROW_LIMIT + 1))


# ── PDF 보안 fetcher (§6.7 — LFI/SSRF 차단) ─────────────────
@pytest.mark.parametrize("url", [
    "file:///etc/passwd",
    "http://169.254.169.254/latest/meta-data",
    "https://example.org/anything",
    "data:text/html,<script>x</script>",
])
def test_fetcher_blocks_external(url):
    with pytest.raises(BlockedUrlError):
        secure_url_fetcher(url)


def test_fetcher_blocks_asset_path_escape():
    with pytest.raises(BlockedUrlError):
        secure_url_fetcher("assets:../../../../etc/passwd")


def test_fetcher_blocks_missing_asset():
    with pytest.raises(BlockedUrlError):
        secure_url_fetcher("assets:no-such-file.png")


# ── 증명서 표기 (발급기관·부서명 — weasyprint 불요, 컨텍스트 캡처) ──
def test_cert_pdf_context_no_placeholder(monkeypatch):
    """표기 결함 회귀 — 발급기관 견본 문구("OO병원장") 금지·부서는 ID 숫자가 아니라 부서명.

    org_name 미설정 시 정직 폴백 "본 의료기관장"(계약서 "본 의료기관" 폴백과 정합).
    """
    from types import SimpleNamespace

    import erp.common.pdf as pdf_mod
    from erp.hr.router import _cert_pdf

    captured = {}

    def _capture(template, ctx):
        captured.update(ctx)
        return b"%PDF-fake"

    monkeypatch.setattr(pdf_mod, "render_pdf", _capture)
    emp = SimpleNamespace(name="김민준", employee_no="E0001", department_id=2,
                          position="팀원", hired_on=date(2020, 1, 1), resigned_on=None)
    issue = SimpleNamespace(id=1, cert_type="EMPLOYMENT", purpose="은행 제출")
    _cert_pdf(emp, issue, dept_name="행정팀")
    assert captured["hospital_name"] == "본 의료기관장"    # org_name 미설정 폴백
    assert "OO병원장" not in str(captured.values())        # 견본 문구 근절
    assert captured["department"] == "행정팀"              # ID(2)가 아니라 부서명
    _cert_pdf(emp, issue, dept_name=None)
    assert captured["department"] == "-"                   # 부서 미지정 표기


# ── PDF 렌더 (weasyprint 런타임 필요 — 미설치 환경 skip) ────
weasyprint = pytest.importorskip("weasyprint")


def test_certificate_pdf_renders():
    pdf = render_pdf("certificate.html", {
        "issue_no": "CERT-000001", "title": "재 직 증 명 서", "kind_text": "재직",
        "name": "김민준", "employee_no": "E0001", "department": "행정팀", "position": "팀원",
        "hired_on": date(2020, 1, 1), "until": "현재", "purpose": "은행 제출",
        "issued_on": "2026-06-07", "hospital_name": "본 의료기관장"})
    assert pdf.startswith(b"%PDF")


def test_contract_pdf_renders():
    pdf = render_pdf("contract.html", {
        "title": "급식 위탁계약", "contract_no": "OC-1", "type_label": "위수탁계약",
        "status": "ACTIVE", "vendor_name": "청정급식", "category": "MEAL",
        "start_on": date(2026, 1, 1), "end_on": date(2026, 12, 31),
        "effective_on": date(2026, 1, 1), "amount": 8_000_000.0, "auto_renew": True,
        "amount_label": "월 금액", "terms": "1조 ... 2조 ...",
        "amendments": [{"change_type": "AMOUNT", "old_value": "8000000",
                        "new_value": "9000000", "effective_on": date(2026, 7, 1)}],
        "org_name": "OO의료재단", "sign_status": "COMPLETED", "generated_on": "2026-06-29"})
    assert pdf.startswith(b"%PDF")


def test_contract_pdf_escapes_injection():
    """autoescape — 계약 자유입력(terms·vendor)이 url() 주입으로 이어지지 않음."""
    pdf = render_pdf("contract.html", {
        "title": "t", "contract_no": "OC-2", "type_label": "구매계약", "status": "DRAFT",
        "vendor_name": '<style>body{background:url("file:///etc/passwd")}</style>',
        "category": "IT", "start_on": date(2026, 1, 1), "end_on": date(2026, 12, 31),
        "effective_on": None, "amount": 100.0, "auto_renew": False,
        "amount_label": "계약 총액", "terms": "<script>alert(1)</script>", "amendments": [],
        "org_name": "OO", "sign_status": None, "generated_on": "2026-06-29"})
    assert pdf.startswith(b"%PDF")


def test_payslip_pdf_renders_and_escapes_injection():
    """Jinja2 autoescape — 악성 입력이 url() 주입으로 이어지지 않음 (§6.7)."""
    pdf = render_pdf("payslip.html", {
        "yyyymm": "2026-06",
        "name": '<style>body{background:url("file:///etc/passwd")}</style>',
        "employee_no": "E0001", "birth_date": "1990-01-01", "position": "간호사",
        "org_name": "위루비병원", "paid_on": "2026-06-25",
        "earnings": [{"item_name": "기본급", "amount": 3_000_000.0, "taxable": True, "formula": None},
                     {"item_name": "연장근로수당", "amount": 300_000.0, "taxable": True,
                      "formula": "10h × 20,000 × 1.5"},
                     {"item_name": "식대", "amount": 200_000.0, "taxable": False, "formula": None}],
        "deductions": [{"item_name": "국민연금", "amount": 144_000.0, "taxable": False,
                        "formula": "3,200,000 × 4.5%"}],
        "total_earning": 3_500_000.0, "total_deduction": 144_000.0, "net": 3_356_000.0,
        "incentive_detail": {"incentive_type": "NET_COLLECTION", "rate_pct": 15,
                             "revenue": 45_000_000.0, "net_revenue": 34_900_000.0,
                             "amount": 5_235_000.0, "formula": "순수입 34,900,000 × 15%"},
        "confirmed_at": "2026-06-07"})
    assert pdf.startswith(b"%PDF")  # escape되어 텍스트로 렌더 — fetch 미발생


# ── 기관 로고(#6) — data:image 인라인 허용, PDF 반영 ─────────
def _png_b64() -> str:
    """유효한 1x1 PNG base64(PIL 생성 — weasyprint/PIL 디코드 보장)."""
    import base64 as _b64

    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGBA", (1, 1), (200, 0, 0, 255)).save(buf, "PNG")
    return _b64.b64encode(buf.getvalue()).decode()


_PNG = _png_b64()


def test_fetcher_allows_inline_image():
    """기관 로고용 data:image 인라인은 허용(네트워크·파일 접근 없음)."""
    out = secure_url_fetcher(f"data:image/png;base64,{_PNG}")
    assert out is not None  # 차단 안 됨(BlockedUrlError 미발생)


def test_fetcher_still_blocks_non_image_data():
    """data:text/html 등 비이미지 data:는 계속 차단(스크립트 인라인 방지)."""
    with pytest.raises(BlockedUrlError):
        secure_url_fetcher("data:text/html,<b>x</b>")


def test_certificate_pdf_embeds_logo():
    """증명서 템플릿에 기관 로고 data:image가 있으면 PDF 렌더(예외 없이 bytes)."""
    from types import SimpleNamespace
    from erp.hr.router import _cert_pdf
    emp = SimpleNamespace(name="김직원", employee_no="E-1", position="간호사",
                          hired_on=date(2022, 1, 1), resigned_on=None)
    issue = SimpleNamespace(id=1, cert_type="EMPLOYMENT", purpose="제출용")
    pdf = _cert_pdf(emp, issue, dept_name="간호부", org_name="위루비병원",
                    logo=f"data:image/png;base64,{_PNG}")
    assert pdf[:4] == b"%PDF"


async def test_org_logo_data_uri(session, tmp_path, monkeypatch):
    """기관 로고 첨부(entity_type=org) → data:image URI. 이미지 없으면 None."""
    from erp.common.attachment_service import AttachmentService
    from erp.common.org_service import org_logo_data_uri
    from erp.config import get_settings
    import base64 as _b64
    monkeypatch.setattr(get_settings(), "attachment_dir", str(tmp_path), raising=False)
    monkeypatch.setattr(get_settings(), "attachment_max_bytes", 5_000_000, raising=False)
    assert await org_logo_data_uri(session) is None            # 없을 때
    await AttachmentService().save(session, entity_type="org", entity_id=1, filename="logo.png",
                                   content_type="image/png", data=_b64.b64decode(_PNG), uploaded_by=1)
    uri = await org_logo_data_uri(session)
    assert uri is not None and uri.startswith("data:image/png;base64,")
