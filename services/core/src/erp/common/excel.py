"""Excel 생성 — openpyxl write_only (design §11.0).

5천 행 이하 동기 응답. 초과분 비동기 잡 전환은 운영 진입 시 (§11.0 결정).
"""
import io
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

SYNC_ROW_LIMIT = 5_000

_HEADER_FILL = PatternFill("solid", fgColor="DDE5F0")
_HEADER_FONT = Font(bold=True, size=10)


def build_xlsx(*, sheet_name: str, headers: list[str],
               rows: list[list[object]]) -> bytes:
    """write_only 모드 — 대량 행 메모리 효율 (design §11.0)."""
    if len(rows) > SYNC_ROW_LIMIT:
        # 비동기 잡 전환 전까지 상한 강제 — 잘림을 숨기지 않음 (no silent caps)
        raise ValueError(f"동기 export 상한({SYNC_ROW_LIMIT}행) 초과: {len(rows)}행")

    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=sheet_name[:31])

    header_cells = []
    for h in headers:
        from openpyxl.cell import WriteOnlyCell
        cell = WriteOnlyCell(ws, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        header_cells.append(cell)
    ws.append(header_cells)

    for row in rows:
        ws.append([_coerce(v) for v in row])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _coerce(v: object) -> object:
    if isinstance(v, datetime):
        # 반출본도 화면과 같은 KST로 — raw timestamptz(UTC aware)를 그대로 쓰면 셀은 +00:00,
        # 같은 데이터의 화면은 +09:00이라 반출 기준 대사·감사에서 9h(날짜 경계엔 1일) 어긋난다.
        from erp.common.kst import kst_iso
        return kst_iso(v)
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, bool):
        return "Y" if v else "N"
    return v
